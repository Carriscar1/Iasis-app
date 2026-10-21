"""
IASIS — uma animação por slide, para apresentar o projeto parte por parte.

    IASIS_LIMPO=1 python hardware/video/gerar_video_v2.py   # vídeo sem legenda nem título
    python hardware/video/gerar_slides.py

Corta o vídeo limpo (saida_slides/IASIS_video_stand_v2.mp4) em um clipe por fase
e monta uma apresentação PowerPoint com o clipe de cada fase, o título, os
tópicos e, nas notas do apresentador, a fala do Iasinho daquela fase.

Saídas em hardware/video/saida_slides/:
  clipes/NN_nome.mp4          um por slide, na ordem da apresentação
  IASIS_apresentacao.pptx     os mesmos slides, com os vídeos dentro
"""
from __future__ import annotations

import subprocess
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Pt

AQUI = Path(__file__).resolve().parent
SAIDA = AQUI / "saida_slides"
VIDEO = SAIDA / "IASIS_video_stand_v2.mp4"

# Início de cada cena no vídeo (s), na ordem de CENAS em gerar_video_v2.py
C = dict(boas_vindas=0, problema=8, solucao=18, por_fora=32, de_cima=46, desmontando=55,
         por_dentro=71, funciona=125, final=147)

# (arquivo, início, fim, título, tópicos, fala do Iasinho nas notas)
SLIDES = [
    ("01_capa", 0, 8, "IASIS", ["Sua saúde em dia", "Dispensador · Pulseira · Aplicativo"],
     "Oi! Eu sou o Iasinho. Vem comigo conhecer o IASIS!"),
    ("02_problema", 8, 18, "O problema",
     ["50% dos pacientes crônicos não seguem o tratamento", "Esquecem ou erram o horário",
      "Fonte: OMS, 2003"],
     "Sabia que metade dos pacientes com doenças crônicas esquece ou erra o horário dos remédios?"),
    ("03_solucao", 18, 32, "A solução",
     ["Dispensador: dose certa na hora certa", "Pulseira: avisa e confirma quem toma",
      "Aplicativo: adesão e aviso ao cuidador"],
     "O IASIS junta três partes: um dispensador, uma pulseira e um aplicativo. "
     "E a dose só sai quando o paciente encosta a pulseira no aparelho."),
    ("04_por_fora", 32, 46, "O dispensador por fora",
     ["Impresso em 3D pela equipe", "Tela, logo com o leitor e bandeja"],
     "Esse é o dispensador do IASIS, impresso em 3D pela nossa equipe. "
     "Na frente: a telinha, a nossa logo e a bandeja onde o remédio chega."),
    ("05_visto_de_cima", 46, 55, "Visto de cima",
     ["Tampa presa por 4 parafusos", "Disco com 8 compartimentos"],
     "Visto de cima: a tampa é presa por quatro parafusos. "
     "Tirando a tampa, aparecem os oito compartimentos do disco."),
    ("06_desmontando", 55, 71, "Desmontando a caixa",
     ["11 peças impressas em 3D", "Motor, servo, leitor e tela", "Tudo volta com encaixes"],
     "Agora, vamos desmontar a caixa! São onze peças impressas, mais o motor, o servo, "
     "o leitor e a tela. E tudo volta pro lugar com encaixes."),
]
DENTRO = [
    ("disco", "Por dentro · O disco", ["8 compartimentos", "Um para cada horário"],
     "Esse é o disco: oito compartimentos, um para cada horário de remédio."),
    ("motor", "Por dentro · O motor", ["Motor de passo NEMA 17", "Gira o disco até a dose da vez"],
     "Embaixo dele, o motor de passo gira o disco certinho até a dose da vez."),
    ("trava", "Por dentro · A trava", ["Servo que tampa a saída", "Só abre depois da pulseira"],
     "Essa é a trava: o servo só abre a saída depois que a pulseira é validada."),
    ("calha", "Por dentro · A calha", ["Leva o remédio até a bandeja"],
     "A calha leva o remédio até a bandeja, direto para a mão do paciente."),
    ("leitor", "Por dentro · O leitor da pulseira", ["Leitor RFID atrás do coração", "É só encostar"],
     "Atrás do coração da logo fica o leitor da pulseira. É só encostar!"),
    ("tela", "Por dentro · A tela", ["LCD 16×2", "Horário e status"],
     "E aqui a tela mostra o horário e o que está acontecendo."),
]
for k, (nome, tit, top, fala) in enumerate(DENTRO):
    a = C["por_dentro"] + 9 * k
    SLIDES.append((f"{7 + k:02d}_dentro_{nome}", a, a + 9, tit, top, fala))
FUNCIONA = [  # fases da cena "Como funciona" (tempos dentro da cena)
    ("hora", 0.0, 3.1, "1. Hora do remédio", ["A tela avisa", "O celular também"],
     "Chegou a hora do remédio: a tela avisa, e o celular também."),
    ("pulseira", 3.1, 6.4, "2. A pulseira", ["Encosta no coração da logo", "Leitura RFID"],
     "O paciente encosta a pulseira no coração da logo..."),
    ("disco", 6.4, 10.1, "3. O disco gira", ["Até o compartimento certo"],
     "...o disco gira até o compartimento certo..."),
    ("trava", 10.1, 12.9, "4. A trava abre", ["Só depois da pulseira validada"], "...a trava abre..."),
    ("dose", 12.9, 17.7, "5. A dose chega", ["Pela calha até a bandeja"],
     "...e a dose escorrega pela calha até a bandeja!"),
    ("registro", 17.7, 22.0, "Dose registrada", ["Fica tudo no app", "O cuidador acompanha"],
     "Pronto! Fica tudo registrado no app, e o cuidador acompanha."),
]
for k, (nome, a, b, tit, top, fala) in enumerate(FUNCIONA):
    SLIDES.append((f"{13 + k:02d}_funciona_{nome}", C["funciona"] + a, C["funciona"] + b,
                   "Como funciona · " + tit, top, fala))
SLIDES.append(("19_obrigado", C["final"], C["final"] + 12, "Obrigado!",
               ["IASIS: sua saúde em dia", "ODS 3 e ODS 9", "Venha testar no stand!"],
               "Esse é o IASIS: sua saúde em dia. Vem testar com a gente aqui no stand!"))


def cortar():
    pasta = SAIDA / "clipes"
    pasta.mkdir(parents=True, exist_ok=True)
    for nome, a, b, *_ in SLIDES:
        destino = pasta / f"{nome}.mp4"
        subprocess.run(["ffmpeg", "-v", "error", "-y", "-ss", f"{a:.2f}", "-to", f"{b:.2f}", "-i", str(VIDEO),
                        "-an", "-c:v", "libx264", "-crf", "20", "-preset", "medium", "-pix_fmt", "yuv420p",
                        "-movflags", "+faststart", str(destino)], check=True)
        # capa do vídeo no slide: o quadro do meio do clipe
        subprocess.run(["ffmpeg", "-v", "error", "-y", "-ss", f"{(b - a) / 2:.2f}", "-i", str(destino),
                        "-frames:v", "1", str(destino.with_suffix(".jpg"))], check=True)
    return pasta


def rgb(h):
    return RGBColor.from_string(h.lstrip("#"))


def apresentacao(pasta):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)    # 16:9
    W, H = prs.slide_width, prs.slide_height
    vazio = prs.slide_layouts[6]
    for nome, a, b, titulo, topicos, fala in SLIDES:
        s = prs.slides.add_slide(vazio)
        fundo = s.background.fill
        fundo.gradient()
        fundo.gradient_angle = 90
        fundo.gradient_stops[0].color.rgb = rgb("#1f1611")
        fundo.gradient_stops[1].color.rgb = rgb("#3a2a20")
        # vídeo à esquerda (60 % da largura, 16:9)
        vw = int(W * 0.60)
        vh = int(vw * 9 / 16)
        x0, y0 = int(W * 0.03), int((H - vh) / 2)
        s.shapes.add_movie(str(pasta / f"{nome}.mp4"), x0, y0, vw, vh,
                           poster_frame_image=str(pasta / f"{nome}.jpg"), mime_type="video/mp4")
        # título e tópicos à direita
        tx = x0 + vw + int(W * 0.03)
        caixa = s.shapes.add_textbox(tx, int(H * 0.22), W - tx - int(W * 0.03), int(H * 0.6))
        tf = caixa.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = titulo
        p.font.size, p.font.bold, p.font.color.rgb = Pt(30), True, rgb("#f6efe6")
        for t in topicos:
            q = tf.add_paragraph()
            q.text = "• " + t
            q.font.size, q.font.color.rgb = Pt(18), rgb("#e8d5b9")
            q.space_before = Pt(10)
        s.notes_slide.notes_text_frame.text = fala
    destino = SAIDA / "IASIS_apresentacao.pptx"
    prs.save(destino)
    return destino


if __name__ == "__main__":
    pasta = cortar()
    print("clipes:", pasta)
    print("pptx:", apresentacao(pasta))
